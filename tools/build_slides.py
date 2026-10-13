"""Build the Round 2 deck from a live evaluation run.

Numbers are written from src.evaluation. They are not typed in by hand.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from src.evaluation import evaluate_suite

NAVY = RGBColor(0x0B, 0x12, 0x20)
INK = RGBColor(0xF4, 0xF1, 0xEA)
MUTED = RGBColor(0xB7, 0xC0, 0xCE)
EMERALD = RGBColor(0x1F, 0x8A, 0x70)
CARD = RGBColor(0x14, 0x1C, 0x2B)


def _fmt(value) -> str:
    if value is None:
        return "undefined"
    try:
        if value != value:  # NaN
            return "undefined"
    except Exception:
        return "undefined"
    return f"{float(value):.3f}"


def _fill(slide, color: RGBColor) -> None:
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _box(slide, text, left, top, width, height, size=18, color=INK, bold=False):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = "Calibri"
    return frame


def _bullets(slide, lines, left, top, width, height, size=16):
    shape = slide.shapes.add_textbox(left, top, width, height)
    frame = shape.text_frame
    frame.word_wrap = True
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_after = Pt(8)
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = INK
        run.font.name = "Calibri"
    return frame


def build(path: Path) -> None:
    report = evaluate_suite()
    detection = report["detection"]
    operations = report["operations"]
    findings = report["findings"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Problem and solution
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "ECHO", Inches(0.6), Inches(0.4), Inches(12), Inches(0.6), 36, EMERALD, True)
    _box(slide, "Every Shortage Leaves a Signal", Inches(0.6), Inches(1.05), Inches(12), Inches(0.4), 20, INK, False)
    _bullets(
        slide,
        [
            "Problem. A fresh-produce counter stocks out after the cause has already moved through growers, sheds, and hubs.",
            "Planners see the empty shelf. They rarely get to compare feasible responses on the same facts.",
            "Solution. A synthetic early-warning and what-if prototype: detect thin cover, trace downstream exposure, and compare do-nothing, alternative sourcing, and redistribution.",
            "Status. Working local prototype. Illustrative Indian place labels. Not a live supply chain and not a field measurement.",
        ],
        Inches(0.6),
        Inches(1.8),
        Inches(12),
        Inches(4.8),
        20,
    )
    _box(slide, "Team Lunaré  ·  Commit for Good Round 2  ·  Synthetic demo", Inches(0.6), Inches(6.8), Inches(12), Inches(0.4), 14, MUTED)

    # 2. Architecture
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Architecture and workflow", Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), 28, EMERALD, True)
    _bullets(
        slide,
        [
            "1. data_generator.py builds a seeded demand plan. It does not step inventory.",
            "2. network.py holds the synthetic graph: suppliers, warehouses, hubs, retailers, lanes, capacities, opening stock.",
            "3. simulation.py steps one day at a time: production, arrivals, assumed spoilage, lost sales, then shipments.",
            "4. risk_engine.py scores that morning only. It cannot see future demand or future stockouts.",
            "5. interventions.py applies do-nothing, alternative sourcing, or redistribution under the same seed.",
            "6. evaluation.py replays held-out scenarios and calls sklearn for precision, recall, and F1.",
            "7. Streamlit (app.py) renders the run. No API keys, database, or runtime network calls.",
        ],
        Inches(0.6),
        Inches(1.2),
        Inches(12),
        Inches(5.5),
        18,
    )

    # 3. Technical approach
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Technical approach and synthetic data", Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), 28, EMERALD, True)
    _bullets(
        slide,
        [
            "Crops: tomato, onion, potato. Unit: crates. Step: one day. Unmet demand is lost, not backlogged.",
            "Alert if stock plus pipeline due before the next primary replenishment cannot cover expected requested demand, or if the fixed-weight score reaches 0.40.",
            "Weights, not fit on the test set: cover 0.50, delay 0.20, demand deviation 0.15, upstream capacity 0.15.",
            "Requested demand is used, never sales. A stockout must not look like a quiet market.",
            "Spoilage is an assumption on hub and retailer on-hand stock only (tomato 2%, onion 0.8%, potato 0.3%).",
            "Baseline comparator: on-hand cover below 2 days. It ignores pipeline, delay, and supplier capacity.",
            "Alternate lanes and lateral moves are attempted only when stock and route capacity exist. Donors keep 1 day of outflow.",
        ],
        Inches(0.6),
        Inches(1.15),
        Inches(12),
        Inches(5.6),
        17,
    )

    # 4. Prototype
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Implemented prototype", Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), 28, EMERALD, True)
    _bullets(
        slide,
        [
            "Runnable Streamlit app: streamlit run app.py.",
            "Sections: executive overview, supply network, early warning, disruption lab, parallel futures, audit, evaluation, data explorer, about.",
            "Controls that change the scenario: preset, seed, horizon, shock type, supplier, route, capacity cut, demand multiplier, delay, window. Applied by Run simulation.",
            "Reset restores the default Nagpur–Delhi transport delay. Compare interventions opens the futures page.",
            "CSV download is available for catalogue tables, shipments, risk rows, and policy summaries.",
            "Screenshots: capture them from a local run. This deck does not embed a fake product shot.",
            "No deployed URL is claimed.",
        ],
        Inches(0.6),
        Inches(1.15),
        Inches(12),
        Inches(5.6),
        18,
    )

    # 5. Evaluation — actual rows
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Evaluation — computed on this build", Inches(0.6), Inches(0.3), Inches(12), Inches(0.45), 28, EMERALD, True)
    _box(
        slide,
        "Held-out synthetic scenarios. Day-level labels use a 3-day window. Undefined metrics are not shown as zero.",
        Inches(0.6),
        Inches(0.85),
        Inches(12),
        Inches(0.4),
        14,
        MUTED,
    )
    lines = []
    for row in detection.itertuples(index=False):
        lines.append(
            f"{row.scenario}: ECHO P {_fmt(row.echo_precision)} R {_fmt(row.echo_recall)} F1 {_fmt(row.echo_f1)} "
            f"| baseline P {_fmt(row.baseline_precision)} R {_fmt(row.baseline_recall)} F1 {_fmt(row.baseline_f1)} "
            f"| false alarms ECHO {int(row.false_alarms_echo)} / baseline {int(row.false_alarms_baseline)}"
        )
    _bullets(slide, lines[:8], Inches(0.45), Inches(1.3), Inches(12.4), Inches(5.6), 13)

    # 6. Feasibility
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Feasibility, beneficiaries, responsible impact", Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), 26, EMERALD, True)
    operating_lines = []
    for scenario_name, group in operations.groupby("scenario", sort=False):
        base = group.loc[group["policy"] == "do_nothing"].iloc[0]
        best = group.sort_values(["unmet_demand_crates", "stockout_events", "transport_cost_units"]).iloc[0]
        operating_lines.append(
            f"{scenario_name}: do-nothing unmet {base.unmet_demand_crates:.0f} crates; "
            f"lowest unmet in this run is {best.policy_label} at {best.unmet_demand_crates:.0f}."
        )
    _bullets(
        slide,
        [
            "Intended future users: replenishment planners, producer organisations, and city market managers. This demo does not serve them with live data.",
            "Actions are advisory. A wrong alert could push stock away from a counter that needed it.",
            *operating_lines[:4],
            "Interventions are allowed to look worse. Read the operating table before claiming a winner.",
        ],
        Inches(0.5),
        Inches(1.1),
        Inches(12.3),
        Inches(5.8),
        15,
    )

    # 7. Future
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Future scope and real-world validation", Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), 28, EMERALD, True)
    _bullets(
        slide,
        [
            "Not in this build: live data, an optimiser, rupee costs, or a claim of reduced food waste.",
            "Validation would need permissioned stock and requested-demand logs, and a lane list operators recognise.",
            "Spoilage percentages need a study. The current rates are assumptions.",
            "A prospective test should log alerts using only past data, then compare them with later stockouts.",
            "Someone must own the operating decision. This repository does not name a deployment partner.",
            "Citations for national food-loss statistics are intentionally omitted until the team verifies them.",
        ],
        Inches(0.6),
        Inches(1.2),
        Inches(12),
        Inches(5.4),
        18,
    )

    # 8. Team
    slide = prs.slides.add_slide(blank)
    _fill(slide, NAVY)
    _box(slide, "Team and links", Inches(0.6), Inches(0.35), Inches(12), Inches(0.5), 28, EMERALD, True)
    _bullets(
        slide,
        [
            "Team: Lunaré",
            "Team leader: Akshadha Saxena",
            "Other members: add only the people who are actually on the team.",
            "GitHub repository: not published from this build. Add the URL after the push.",
            "Live demo: not deployed from this build. Add the URL only after Streamlit Community Cloud succeeds.",
            "Local demo command: streamlit run app.py",
            "Tests: pytest -q",
        ],
        Inches(0.6),
        Inches(1.3),
        Inches(12),
        Inches(5),
        20,
    )

    # Keep a text companion so the numbers can be checked without opening PowerPoint.
    notes = path.with_suffix(".findings.txt")
    notes.write_text("\n".join(findings), encoding="utf-8")
    prs.save(path)


if __name__ == "__main__":
    destination = Path(__file__).resolve().parents[1] / "docs" / "ECHO_round2.pptx"
    build(destination)
    print(destination)
